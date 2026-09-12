"""Local tools the agent can call. 100% on-premise: no external network calls."""
import base64
import datetime
import io
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image

from . import config
from .rag import search as rag_search
from .llm import client


def _resolve(path: str) -> Path:
    """Confine every file operation to the project folder."""
    p = Path(str(path).strip().strip('"'))
    if not p.is_absolute():
        p = config.BASE_DIR / p
    p = p.resolve()
    if p != config.BASE_DIR and config.BASE_DIR not in p.parents:
        raise ValueError(f"access outside the project folder is not allowed: {p}")
    return p


resolve_path = _resolve  # public alias used by the agent's policy guards


# ---------------- file tools ----------------

def tool_list_files(path: str = ".") -> str:
    p = _resolve(path)
    if not p.exists():
        return f"ERROR: no such path: {path}"
    if p.is_file():
        return p.name
    rows = [("dir  " if c.is_dir() else "file ") + c.name for c in sorted(p.iterdir())]
    return "\n".join(rows) if rows else "(empty folder)"


def tool_read_file(path: str) -> str:
    p = _resolve(path)
    if not p.is_file():
        return f"ERROR: no such file: {path}"
    text = p.read_text(encoding="utf-8", errors="replace")
    return text[:6000] + ("\n...[truncated]" if len(text) > 6000 else "")


def tool_read_pdf(path: str) -> str:
    import pymupdf

    p = _resolve(path)
    if not p.is_file():
        return f"ERROR: no such file: {path}"
    document = pymupdf.open(str(p))
    pages = [f"--- Page {number} ---\n{text}" for number, page in enumerate(document, 1)
             if (text := page.get_text().strip())]
    document.close()
    result = "\n\n".join(pages) or "(no selectable text found; use ocr_image for a scanned PDF)"
    return result[:12000] + ("\n...[truncated]" if len(result) > 12000 else "")


def tool_read_excel(path: str, sheet_name: str = "") -> str:
    from openpyxl import load_workbook

    p = _resolve(path)
    if not p.is_file():
        return f"ERROR: no such file: {path}"
    workbook = load_workbook(str(p), read_only=True, data_only=True)
    sheets = [sheet_name] if sheet_name else workbook.sheetnames
    sections = []
    for name in sheets:
        if name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                rows.append("\t".join(values))
        sections.append(f"--- Sheet: {name} ---\n" + ("\n".join(rows) or "(empty)"))
    workbook.close()
    if not sections:
        return f"ERROR: no matching worksheet: {sheet_name}"
    result = "\n\n".join(sections)
    return result[:12000] + ("\n...[truncated]" if len(result) > 12000 else "")


def tool_read_docx(path: str) -> str:
    from docx import Document

    p = _resolve(path)
    if not p.is_file():
        return f"ERROR: no such file: {path}"
    document = Document(str(p))
    sections = [paragraph.text.strip() for paragraph in document.paragraphs
                if paragraph.text.strip()]
    for number, table in enumerate(document.tables, 1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        sections.append(f"--- Table {number} ---\n" + "\n".join(rows))
    result = "\n".join(sections) or "(document is empty)"
    return result[:12000] + ("\n...[truncated]" if len(result) > 12000 else "")


def tool_read_pptx(path: str) -> str:
    from pptx import Presentation

    p = _resolve(path)
    if not p.is_file():
        return f"ERROR: no such file: {path}"
    presentation = Presentation(str(p))
    slides = []
    for number, slide in enumerate(presentation.slides, 1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                rows = ["\t".join(cell.text.strip() for cell in row.cells)
                        for row in shape.table.rows]
                parts.append("[Table]\n" + "\n".join(rows))
        slides.append(f"--- Slide {number} ---\n" + ("\n".join(parts) or "(no text)"))
    result = "\n\n".join(slides) or "(presentation is empty)"
    return result[:16000] + ("\n...[truncated]" if len(result) > 16000 else "")


def tool_write_file(path: str, content: str) -> str:
    p = _resolve(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return f"OK: wrote {len(content)} characters to {p.relative_to(config.BASE_DIR)}"


def tool_edit_file(path: str, search: str, replace: str) -> str:
    """Surgical edit: replace exactly one occurrence of `search` in the file.
    Preferred over write_file for small fixes - no need to re-emit the file."""
    p = _resolve(path)
    if not p.is_file():
        return f"ERROR: no such file: {path}"
    text = p.read_text(encoding="utf-8", errors="ignore")
    note = ""
    n = text.count(search)
    if (n == 0 and len(search) > 2 and search[0] == search[-1]
            and search[0] in "'\"`" and text.count(search[1:-1]) > 0):
        # common small-model mistake: wrapping the target word in quotes
        search = search[1:-1]
        n = text.count(search)
        if len(replace) > 2 and replace[0] == replace[-1] and replace[0] in "'\"`":
            replace = replace[1:-1]
        note = " (NOTE: your search/replace were wrapped in quotes; unquoted text was used.)"
    if n == 0:
        return ("ERROR: the search text was not found in the file. Copy the text to change "
                "VERBATIM from the read_file output, exactly as it appears in the file "
                "(e.g. readings.lenght - do NOT add quotes around it).")
    if n > 1:
        return (f"ERROR: the search text occurs {n} times. Include more surrounding "
                "lines so it matches exactly once.")
    p.write_text(text.replace(search, replace, 1), encoding="utf-8")
    return f"OK: replaced 1 occurrence in {p.relative_to(config.BASE_DIR)}.{note}"


# ---------------- code execution sandbox ----------------

def _format_proc(proc) -> str:
    out = f"exit_code={proc.returncode}"
    if proc.stdout and proc.stdout.strip():
        out += "\nSTDOUT:\n" + proc.stdout.strip()[:4000]
    if proc.stderr and proc.stderr.strip():
        out += "\nSTDERR:\n" + proc.stderr.strip()[:4000]
    if proc.returncode == 0 and not (proc.stdout or "").strip():
        out += " (no output)"
    return out


def tool_run_python(code: str) -> str:
    script = config.SANDBOX_DIR / "_agent_run.py"
    script.write_text(code, encoding="utf-8")
    try:
        proc = subprocess.run([sys.executable, str(script)],
                              cwd=str(config.SANDBOX_DIR), capture_output=True,
                              text=True, timeout=config.CODE_TIMEOUT)
    except subprocess.TimeoutExpired:
        return f"ERROR: execution timed out after {config.CODE_TIMEOUT}s"
    return _format_proc(proc)


def tool_run_command(command: str) -> str:
    # A foreground server never exits, so subprocess.run would consume the
    # entire command budget and prevent the worker from receiving an observation.
    if re.search(r"(?:http\.server|vite|webpack|live-server|flask run|uvicorn)",
                 command, re.IGNORECASE):
        return ("ERROR: this is a long-running development server and cannot be run "
                "in the foreground by run_command. For a static HTML/CSS task, "
                "inspect the files directly and finish; do not start a server. "
                "If a server is explicitly required, start it outside the agent run.")
    try:
        proc = subprocess.run(command, shell=True, cwd=str(config.BASE_DIR),
                              capture_output=True, text=True,
                              timeout=config.CODE_TIMEOUT)
    except subprocess.TimeoutExpired:
        return f"ERROR: command timed out after {config.CODE_TIMEOUT}s"
    return _format_proc(proc)


# ---------------- vision / OCR ----------------

IMAGE_MIMES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".gif": "image/gif",
}
IMAGE_EXTENSIONS = frozenset(IMAGE_MIMES)

_ocr_engine = None


def _get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        from paddleocr import PaddleOCR
        _ocr_engine = PaddleOCR(use_angle_cls=True, lang="en",
                                use_gpu=False, enable_mkldnn=False)
    return _ocr_engine


def tool_ocr_image(path: str) -> str:
    p = _resolve(path)
    if not p.is_file():
        return f"ERROR: no such file: {path}"
    results = _get_ocr().ocr(str(p), cls=True)
    lines = [line[1][0] for line in results[0]] if results and results[0] else []
    return "OCR text extracted:\n" + ("\n".join(lines) if lines else "(no text detected)")


def load_image_as_png_b64(path: Path) -> str:
    """Normalize supported image formats before sending them to the vision model."""
    with Image.open(path) as image:
        rgb_image = image.convert("RGB")
        buffer = io.BytesIO()
        rgb_image.save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def tool_analyze_image(path: str, question: str = "Describe this image in detail.") -> str:
    p = _resolve(path)
    if not p.is_file():
        return f"ERROR: no such file: {path}"
    mime = IMAGE_MIMES.get(p.suffix.lower())
    if mime is None:
        readers = {
            ".pdf": (tool_read_pdf, "read_pdf"),
            ".docx": (tool_read_docx, "read_docx"),
            ".pptx": (tool_read_pptx, "read_pptx"),
            ".xlsx": (tool_read_excel, "read_excel"),
            ".xlsm": (tool_read_excel, "read_excel"),
        }
        reader, reader_name = readers.get(p.suffix.lower(), (None, "read_file"))
        if reader is not None:
            return (f"ROUTED: analyze_image is not valid for {p.suffix.lower()}; "
                    f"used {reader_name} instead.\n{reader(path)}")
        return (f"ERROR: analyze_image accepts image files, not {p.suffix.lower() or 'this file type'}. "
                f"Use {reader_name}(path) instead.")
    b64 = load_image_as_png_b64(p)
    resp = client.chat.completions.create(
        model=config.MODEL_VISION,
        messages=[{"role": "user", "content": [
            {"type": "text", "text": question},
            {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}},
        ]}],
        temperature=0.1,
        max_tokens=config.VISION_MAX_TOKENS,
        timeout=config.VISION_TIMEOUT,
    )
    return resp.choices[0].message.content


# ---------------- document output ----------------

def tool_create_docx(filename: str, title: str, body: str) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    head = doc.add_heading(title, level=1)
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.add_run(f"Date: {datetime.date.today():%d %B %Y} | Prepared by: Sovereign AI "
                 "Workbench (auto-draft - pending engineer sign-off)").italic = True
    for line in body.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=2)
        elif s.startswith(("- ", "* ")):
            doc.add_paragraph(s[2:], style="List Bullet")
        else:
            doc.add_paragraph(s)
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.add_run("Drafted locally by the AI agent. No data left the premises. "
                 "Requires human engineer sign-off.").italic = True
    out = config.OUTPUTS_DIR / (Path(filename).stem + ".docx")
    doc.save(str(out))
    return f"OK: report saved as outputs/{out.name} (downloadable deliverable)"


def tool_create_pdf(filename: str, title: str, body: str) -> str:
    import pymupdf as fitz

    document = fitz.open()
    page_width, page_height = fitz.paper_size("a4")
    margin = 54
    title_rect = fitz.Rect(margin, margin, page_width - margin, margin + 36)
    body_rect = fitz.Rect(margin, margin + 54, page_width - margin, page_height - margin)

    def add_page() -> fitz.Page:
        page = document.new_page(width=page_width, height=page_height)
        page.insert_textbox(title_rect, title, fontsize=18, fontname="hebo",
                            color=(0.08, 0.16, 0.25))
        return page

    page = add_page()
    cursor = body_rect.y0
    for line in body.splitlines():
        text = line.strip()
        if not text:
            cursor += 8
            continue
        if text.startswith("## "):
            fontsize, color, text = 13, (0.08, 0.16, 0.25), text[3:]
        elif text.startswith("# "):
            fontsize, color, text = 15, (0.08, 0.16, 0.25), text[2:]
        else:
            fontsize, color = 10, (0, 0, 0)
            if text.startswith(("- ", "* ")):
                text = "- " + text[2:]
        box = fitz.Rect(body_rect.x0, cursor, body_rect.x1, body_rect.y1)
        result = page.insert_textbox(box, text, fontsize=fontsize, fontname="helv",
                                     color=color, lineheight=1.25)
        if result < 0:
            page = add_page()
            cursor = body_rect.y0
            page.insert_textbox(fitz.Rect(body_rect.x0, cursor, body_rect.x1, body_rect.y1),
                                text, fontsize=fontsize, fontname="helv", color=color,
                                lineheight=1.25)
            cursor += fontsize * 1.25 * 2
        else:
            cursor = body_rect.y1 - result
            if cursor > body_rect.y1 - 18:
                page = add_page()
                cursor = body_rect.y0
    out = config.OUTPUTS_DIR / (Path(filename).stem + ".pdf")
    document.save(str(out))
    document.close()
    return f"OK: PDF report saved as outputs/{out.name} (downloadable deliverable)"


def tool_create_excel(filename: str, sheet_name: str = "Sheet1", headers: list | None = None,
                      rows: list | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    headers = headers or []
    rows = rows or []
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = str(sheet_name)[:31] or "Sheet1"
    if headers:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(row if isinstance(row, list) else [row])
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[column[0].column_letter].width = min(width, 50)
    out = config.OUTPUTS_DIR / (Path(filename).stem + ".xlsx")
    workbook.save(str(out))
    return f"OK: Excel workbook saved as outputs/{out.name} (downloadable deliverable)"


def tool_create_pptx(filename: str, title: str, slides: list) -> str:
    from pptx import Presentation
    from pptx.util import Pt

    if not isinstance(slides, list) or not slides:
        return "ERROR: slides must be a non-empty list of slide objects"

    presentation = Presentation()
    content_layout = presentation.slide_layouts[1]
    presentation.core_properties.title = str(title)

    for index, slide_data in enumerate(slides, 1):
        if not isinstance(slide_data, dict):
            return f"ERROR: slide {index} must be an object with title and body"
        slide = presentation.slides.add_slide(content_layout)
        slide.shapes.title.text = str(slide_data.get("title", f"Slide {index}"))
        body = slide_data.get("body", slide_data.get("bullets", []))
        lines = body.splitlines() if isinstance(body, str) else body
        if not isinstance(lines, list):
            lines = [str(lines)]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for line_number, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if line_number == 0 else frame.add_paragraph()
            text = str(line).strip()
            paragraph.text = text[2:] if text.startswith(("- ", "* ")) else text
            paragraph.level = 0
            paragraph.font.size = Pt(22)

    out = config.OUTPUTS_DIR / (Path(filename).stem + ".pptx")
    presentation.save(str(out))
    return (f"OK: PowerPoint presentation saved as outputs/{out.name} "
            f"({len(slides)} slides; downloadable deliverable)")


def tool_search_docs(query: str) -> str:
    return rag_search(query)


TOOLS = {
    "list_files":    {"fn": lambda a: tool_list_files(**{"path": ".", **a}), "sig": 'list_files(path=".")'},
    "read_file":     {"fn": lambda a: tool_read_file(**a),   "sig": "read_file(path)"},
    "read_pdf":      {"fn": lambda a: tool_read_pdf(**a),    "sig": "read_pdf(path)"},
    "read_excel":    {"fn": lambda a: tool_read_excel(**a),  "sig": "read_excel(path, sheet_name=\"\")"},
    "read_docx":     {"fn": lambda a: tool_read_docx(**a),   "sig": "read_docx(path)"},
    "read_pptx":     {"fn": lambda a: tool_read_pptx(**a),   "sig": "read_pptx(path)"},
    "write_file":    {"fn": lambda a: tool_write_file(**a),  "sig": "write_file(path, content)"},
    "edit_file":     {"fn": lambda a: tool_edit_file(**a),   "sig": "edit_file(path, search, replace)"},
    "run_python":    {"fn": lambda a: tool_run_python(**a),  "sig": "run_python(code)"},
    "run_command":   {"fn": lambda a: tool_run_command(**a), "sig": "run_command(command)"},
    "ocr_image":     {"fn": lambda a: tool_ocr_image(**a),   "sig": "ocr_image(path)"},
    "analyze_image": {"fn": lambda a: tool_analyze_image(**a), "sig": "analyze_image(path, question=\"Describe this image in detail.\")"},
    "create_docx":   {"fn": lambda a: tool_create_docx(**a), "sig": "create_docx(filename, title, body)"},
    "create_pdf":    {"fn": lambda a: tool_create_pdf(**a), "sig": "create_pdf(filename, title, body)"},
    "create_excel":  {"fn": lambda a: tool_create_excel(**a), "sig": "create_excel(filename, sheet_name, headers, rows)"},
    "create_pptx":   {"fn": lambda a: tool_create_pptx(**a), "sig": "create_pptx(filename, title, slides)"},
    "search_docs":   {"fn": lambda a: tool_search_docs(**a), "sig": "search_docs(query)"},
}

TOOL_MANIFEST = """\
- list_files(path=".") : list files and folders (paths are relative to the project root).
- read_file(path) : read a text file and return its contents.
- read_pdf(path) : extract selectable text from a PDF. For scanned PDFs, use ocr_image(path).
- read_excel(path, sheet_name="") : read one worksheet, or all worksheets when sheet_name is empty.
- read_docx(path) : extract Word paragraphs and table contents.
- read_pptx(path) : extract text and tables from every PowerPoint slide.
- write_file(path, content) : create/overwrite a text file. Code goes to sandbox/ , reports to outputs/ .
- edit_file(path, search, replace) : replace ONE exact text occurrence in an existing file - PREFERRED for small fixes; 'search' must match the file exactly once (copy it from read_file output).
- run_python(code) : run Python 3 code in the local sandbox; returns exit_code, STDOUT, STDERR. Use it to test code.
- run_command(command) : run a shell command FROM THE PROJECT ROOT, with project-relative paths exactly like write_file (examples: "python sandbox/pressure_check.py", "python -m py_compile sandbox/app.py", "cd sandbox\myapp && npx tsc --noEmit", "cd sandbox\myapp && npm run lint"). 60 s timeout.
- ocr_image(path) : extract all text from an image or scanned document using the local OCR engine.
- analyze_image(path, question="Describe this image in detail.") : ask the local vision model a specific question about an image / drawing / photo.
- create_docx(filename, title, body) : create a Word report in outputs/. In body, "## Text" makes a heading, "- text" makes a bullet.
- create_pdf(filename, title, body) : create a PDF report in outputs/. In body, "## Text" makes a heading, "- text" makes a bullet.
- create_excel(filename, sheet_name, headers, rows) : create an Excel workbook in outputs/. Pass headers as a list and rows as a list of lists.
- create_pptx(filename, title, slides) : create a PowerPoint presentation in outputs/. Pass exactly one object per requested slide, like {"title": "...", "body": ["bullet 1", "bullet 2"]}; the title is stored as presentation metadata.
- search_docs(query) : semantic search over ingested reference/SOP documents (RAG); returns relevant passages with source names."""
