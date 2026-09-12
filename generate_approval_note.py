from openai import OpenAI
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

client = OpenAI(api_key="sk-anything", base_url="http://localhost:4000")

def load_context():
    with open("combined_context.txt", "r", encoding="utf-8") as f:
        return f.read()

def generate_approval_note_content(context):
    prompt = f"""You are drafting an internal engineering approval note for a refinery process diagram review.

Grounded context (from OCR + vision analysis of the diagram):
{context}

Write the approval note content in exactly this structure, using these exact headers:

FINDINGS:
(3-5 bullet points on what the diagram shows — key equipment, process flow, any notable elements)

OBSERVATIONS:
(2-3 bullet points on things worth flagging for review — safety-relevant components, unclear areas, or anything an engineer should double check)

RECOMMENDATION:
(1-2 sentences: recommend approval, approval with conditions, or further review — and why)

Do not add any other headers or commentary outside this structure."""

    response = client.chat.completions.create(
        model="reasoning",
        messages=[{"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content

def parse_sections(raw_text):
    sections = {"FINDINGS": "", "OBSERVATIONS": "", "RECOMMENDATION": ""}
    current = None
    for line in raw_text.splitlines():
        stripped = line.strip()
        matched = False
        for key in sections:
            if stripped.upper().startswith(key + ":"):
                current = key
                matched = True
                break
        if not matched and current:
            sections[current] += line + "\n"
    return sections

def build_docx(sections, output_path="Approval_Note.docx"):
    doc = Document()

    title = doc.add_heading("Engineering Approval Note", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run(f"Date: {datetime.date.today().strftime('%d %B %Y')}\n").bold = False
    meta.add_run("Subject: Review of Process Diagram (Pervaporation Membrane System)\n")
    meta.add_run("Prepared by: AI Workbench (Auto-drafted — pending engineer review)")

    doc.add_heading("Findings", level=2)
    for line in sections["FINDINGS"].strip().splitlines():
        line = line.strip("-• ").strip()
        if line:
            doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Observations", level=2)
    for line in sections["OBSERVATIONS"].strip().splitlines():
        line = line.strip("-• ").strip()
        if line:
            doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Recommendation", level=2)
    doc.add_paragraph(sections["RECOMMENDATION"].strip())

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run("Note: This document was drafted by a local AI agent grounded in OCR and vision-model analysis of the source diagram. It requires human engineer sign-off before being treated as an approved record.").italic = True

    doc.save(output_path)
    print(f"Saved: {output_path}")

if __name__ == "__main__":
    context = load_context()
    print("Generating approval note content...")
    raw_content = generate_approval_note_content(context)
    print(f"\n--- Raw model output ---\n{raw_content}\n")

    sections = parse_sections(raw_content)
    build_docx(sections)