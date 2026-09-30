import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    return heading

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p

def main():
    doc = Document()
    
    # Title
    title = doc.add_heading('MRPL Agentic Workbench - End-to-End Architecture Document', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 1. Overview & End-to-End Architecture
    add_heading(doc, '1. Overview & End-to-End Architecture', level=1)
    doc.add_paragraph(
        "The Sovereign Offline Agent Workbench is an on-premise, autonomous AI agent designed "
        "to run fully offline without external network calls. It follows a think->act->observe "
        "loop implemented with LangGraph. The system consists of an Orchestrator/Router, "
        "specialized Worker Models (Reasoning, Coder, Vision), a Sandboxed execution environment, "
        "and an integrated Toolset. It is built to safely execute code, process documents, and "
        "answer complex reasoning tasks while maintaining high reliability through robust guardrails."
    )

    # 2. Tech Stack
    add_heading(doc, '2. Tech Stack', level=1)
    p = doc.add_paragraph()
    p.add_run('Core Orchestration: ').bold = True
    p.add_run('LangGraph (StateGraph), Python.\n')
    p.add_run('LLM Interface: ').bold = True
    p.add_run('LiteLLM (proxying to local Ollama), OpenAI Python Client.\n')
    p.add_run('Document Processing: ').bold = True
    p.add_run('Pillow (Images), python-pptx, python-docx, openpyxl, pymupdf.\n')
    p.add_run('OCR & Vision: ').bold = True
    p.add_run('PaddleOCR, Ollama Vision Models.\n')
    p.add_run('Execution: ').bold = True
    p.add_run('Python subprocess (Sandboxing).\n')

    # 3. Models Used
    add_heading(doc, '3. Models Used (via local Ollama)', level=1)
    doc.add_paragraph("The system intelligently routes tasks to specialized models, swapping them in and out of VRAM:", style='List Bullet')
    doc.add_paragraph("Router & Reasoning: qwen2.5:7b (Handles classification, summaries, document creation)", style='List Bullet')
    doc.add_paragraph("Coder: qwen2.5-coder:7b (Handles writing, testing, and debugging code)", style='List Bullet')
    doc.add_paragraph("Vision: granite3.2-vision:latest / minicpm-v:latest (Handles image analysis)", style='List Bullet')
    doc.add_paragraph("Embeddings: nomic-embed-text (For RAG and semantic search over docs)", style='List Bullet')

    # 4. How Routing Works
    add_heading(doc, '4. How Routing Works', level=1)
    doc.add_paragraph(
        "Routing is handled at the start of the LangGraph execution by the 'classify' node. "
        "The router uses a few-shot prompt to classify the user's task into three categories: 'coding', 'vision', or 'reasoning'. "
        "Depending on the classification, the task is routed to the corresponding worker model. "
        "To optimize VRAM on limited GPUs, the system dynamically unloads models (via `ollama stop`) before loading the next one. "
        "Additionally, there is an override mechanism: if a task is classified as 'coding' but involves creating structured documents (e.g., word, excel, pdf), "
        "it is overridden to 'reasoning' because the coder model struggles with large structured JSON."
    )

    # 5. Orchestration
    add_heading(doc, '5. Orchestration Loop', level=1)
    doc.add_paragraph(
        "The core logic is driven by a LangGraph StateGraph. The state object tracks the task, transcript, history, "
        "pending actions, files accessed, verification status, and step counts. "
        "The loop executes as follows: classify -> agent_step (generates thought/action) -> tool_exec (executes tools and returns observation) -> agent_step. "
        "The loop terminates when the model issues a 'final_answer' or when an orchestrator guardrail forcefully ends the run."
    )

    # 6. Guardrails & Edge Cases
    add_heading(doc, '6. Guardrails & Edge Case Handling', level=1)
    doc.add_paragraph("The system implements strict deterministic guardrails to prevent agentic failure modes:")
    doc.add_paragraph("JSON Repair: Small models often emit invalid JSON (unclosed brackets, unescaped quotes/newlines). The orchestrator includes custom parsers (_repair_doc_action_json, _fix_json_brackets) to fix malformed outputs before crashing.", style='List Bullet')
    doc.add_paragraph("Stuck-Loop Block: If the agent repeats the exact same tool execution 3 times, the orchestrator injects a warning block. At 4 repetitions, it forcefully terminates the run.", style='List Bullet')
    doc.add_paragraph("Verification Guard (Coding): If the agent writes code, it is blocked from emitting a 'final_answer' until it successfully tests the code (exit_code=0).", style='List Bullet')
    doc.add_paragraph("Auto-Finish: If the agent successfully verifies code but keeps inspecting it pointlessly for 2 more steps, the orchestrator forcefully completes the task.", style='List Bullet')
    doc.add_paragraph("File Overwrite Guard: The agent is prevented from overwriting a file (write_file) unless it has read it (read_file) during the current session.", style='List Bullet')
    doc.add_paragraph("Blank Output Streak: If the model outputs empty text 3 consecutive times, the run is terminated.", style='List Bullet')
    doc.add_paragraph("Step Limit: A hard cap of 19 steps prevents infinite agent runs.", style='List Bullet')
    doc.add_paragraph("Timeouts: Strict timeouts are enforced on code execution (60s) and LLM inference (180s). Foreground servers (e.g., http.server) are regex-blocked.", style='List Bullet')

    # 7. Future Improvements
    add_heading(doc, '7. Future Improvements', level=1)
    doc.add_paragraph("More robust sandboxing: Transitioning from standard subprocess isolation to Docker-based containers.", style='List Bullet')
    doc.add_paragraph("Multi-Agent Collaboration: Adding dedicated reviewer agents for code validation.", style='List Bullet')
    doc.add_paragraph("Model Upgrades: Scaling up to larger parameter models when more VRAM becomes available.", style='List Bullet')
    doc.add_paragraph("Streaming UI: Implementing token-by-token streaming for real-time frontend observability.", style='List Bullet')

    # Save
    out_path = os.path.join(r"c:\Users\mamga\OneDrive\Desktop\mrpl-agentic-workbench", "Architecture_Documentation.docx")
    doc.save(out_path)
    print(f"Document saved to {out_path}")

if __name__ == '__main__':
    main()
